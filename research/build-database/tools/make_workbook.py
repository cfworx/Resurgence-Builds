#!/usr/bin/env python3
"""Build the Build-Making Database workbook (one .xlsx, a tab per database section) from this folder's
JSON and CSV files, so the spreadsheet and the repo records never drift apart.

  python3 make_workbook.py OUT.xlsx

Counts (how many weapons can roll a talent, how many talents a weapon can roll, and so on) are
COUNTIF formulas over the Source tabs, so they stay right if a row is edited. Run
scripts/recalc.py from the xlsx skill (or open the file in Excel) to fill in their values.
"""

import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='2B3640')
HEAD_FONT = Font(name=FONT, size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color='555555')
THIN = Side(style='thin', color='D0D4D8')
RARITY_FILL = {'High-End': 'FCE7B8', 'Superior': 'E7D8F7', 'Specialized': 'D6E6FA'}
LONG = {'Talent', 'Talent Text', '2 Set Pieces', '3 Set Pieces', 'Description', 'Note', '2-Piece', '3-Piece', '4-Piece',
        'Gear Sets Note', 'What it is', 'Reading'}


def load(name):
    return json.load(open(os.path.join(BASE, name)))


def rows_csv(name):
    return list(csv.DictReader(open(os.path.join(BASE, name), newline='')))


def sheet(wb, title, headers, rows, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NOTE_FONT
        r0 = 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=headers[c - 1] in LONG)
            cell.border = Border(bottom=THIN)
    for c, h in enumerate(headers, 1):
        if widths and h in widths:
            w = widths[h]
        elif h in LONG:
            w = 60
        else:
            longest = max([len(str(h))] + [len(str(r[c - 1])) for r in rows if c - 1 < len(r) and r[c - 1] is not None][:400])
            w = min(42, max(9, longest + 2))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}'
    ws.row_dimensions[r0].height = 30
    return ws, r0


def shade_rarities(ws, r0, col, n):
    """Colour a Rarity cell the way the game colours the tiles (gold, purple, blue)."""
    for i in range(r0 + 1, r0 + 1 + n):
        cell = ws.cell(row=i, column=col)
        first = str(cell.value or '').split(' / ')[0].split(',')[0].strip()
        if first in RARITY_FILL:
            cell.fill = PatternFill('solid', fgColor=RARITY_FILL[first])


def secs(r):
    return f"{r.get('Start (s)', '')} to {r.get('End (s)', '')} s" if r.get('Start (s)') else ''


def main(out):
    wb = Workbook()
    readme = wb.active
    readme.title = 'Read Me'

    talents = load('weapon-talents.json')
    order = {t['name']: t['listOrder'] for t in talents}
    sources = load('weapon-talent-sources.json')
    arm_w = load('armory-weapons.json')
    arm_g = load('armory-gear.json')
    exotics = load('exotic-weapons.json')
    gs_src = load('gear-set-sources.json')
    ba_src = load('body-armor-talent-sources.json')
    bp_src = load('backpack-talent-sources.json')

    # ---- OS Protocols
    os_rows = [[r['Name'], r['Tab'], r['Rarity'], r['Main Stat'], r['Main Value'], r['Attribute 1'], r['Value 1'],
                r['Attribute 2'], r['Value 2'], r['Attribute 3'], r['Value 3'], r['Talent'], r['Talent Cooldown'],
                r['Shown At'], r['Recording'], secs(r)] for r in rows_csv('os-protocols.csv')]
    ws, r0 = sheet(wb, 'OS Protocols', ['Name', 'Tab', 'Rarity', 'Main Stat', 'Main Value', 'Attribute 1', 'Value 1',
                                         'Attribute 2', 'Value 2', 'Attribute 3', 'Value 3', 'Talent', 'Talent Cooldown',
                                         'Shown At', 'Recording', 'Seconds'], os_rows)
    shade_rarities(ws, r0, 3, len(os_rows))

    # ---- Exotic Weapons
    ex_rows = []
    for e in exotics:
        a = e['weaponAttributes']
        ex_rows.append([e['name'], e['class'], e['shownAt'], e['damageType'], e['damage'], e['rpm'], e['magazine'],
                        e['talentName'], e['talentText'], a['accuracy'], a['reloadSpeed'], a['optimalRange'],
                        a['movementSpeed'], a['stability'], e['modSlots'], e['tamSlots'], len(e['possibleTalents']),
                        '; '.join(e['possibleTalents']), e['description'], e['source']['recording'],
                        f"{e['source']['seconds'][0]} to {e['source']['seconds'][1]} s"])
    sheet(wb, 'Exotic Weapons', ['Name', 'Class', 'Shown At', 'Damage Type', 'Damage', 'RPM', 'Magazine', 'Talent',
                                 'Talent Text', 'Accuracy', 'Reload Speed', 'Optimal Range', 'Movement Speed', 'Stability',
                                 'Mod Slots', 'TAM Slots', 'Possible Talents', 'Possible Talent Names', 'Description',
                                 'Recording', 'Seconds'], ex_rows, widths={'Possible Talent Names': 60})

    # ---- Skill Mod Combo Sets
    smc = [[r['Name'], r['Tab'], int(r['List Order']), r['Skill'], r['2 Set Pieces'], r['3 Set Pieces'],
            r['Talent Cooldown'], r['Rarities'], r['Note'], r['Recording'], secs(r)] for r in rows_csv('skill-mod-combos.csv')]
    sheet(wb, 'Skill Mod Combo Sets', ['Name', 'Tab', 'List Order', 'Skill', '2 Set Pieces', '3 Set Pieces',
                                       'Talent Cooldown', 'Rarities', 'Note', 'Recording', 'Seconds'], smc)

    # ---- Weapon Talents (+ how many weapons can roll each, as a formula over the Sources tab)
    exo_of = {e['talentName']: e['name'] for e in exotics}
    wt_rows = []
    for r in rows_csv('weapon-talents.csv'):
        wt_rows.append([r['Name'], int(r['List Order']), 'yes' if r['Exotic'] == 'yes' else '', r['Talent'],
                        r['Talent Cooldown'], r['Shown At'], None, None, r['Note'], r['Recording'], secs(r)])
    ws, r0 = sheet(wb, 'Weapon Talents', ['Name', 'List Order', 'Exotic', 'Talent', 'Talent Cooldown', 'Shown At',
                                          'Weapons That Can Roll It', 'Rolls On', 'Note', 'Recording', 'Seconds'], wt_rows)
    for i, r in enumerate(wt_rows, r0 + 1):
        if r[2] == 'yes':
            ws.cell(row=i, column=7, value=1)
            ws.cell(row=i, column=8, value=f'{exo_of.get(r[0], "its exotic")} only (the exotic\'s own talent)')
        else:
            ws.cell(row=i, column=7, value=f"=COUNTIF('Weapon Talent Sources'!$A:$A,A{i})")
            t = next(x for x in sources if x['talent'] == r[0])
            ws.cell(row=i, column=8, value=', '.join(t['classOrder']))
        for c in (7, 8):
            ws.cell(row=i, column=c).font = BODY_FONT

    # ---- Weapon Talent Sources: one row per talent and weapon, in the order the Source panel lists them
    src_rows = []
    for t in sorted(sources, key=lambda t: order.get(t['talent'], 999)):
        for w in t['weapons']:
            src_rows.append([t['talent'], w['name'], w['class'], w['armoryTier'], ' / '.join(w['rarities']),
                             'weapon card' if w['match'] == 'confirmed' else 'icon',
                             t['source']['recording'], f"{t['source']['seconds'][0]} to {t['source']['seconds'][1]} s"])
    ws, r0 = sheet(wb, 'Weapon Talent Sources', ['Talent', 'Weapon', 'Class', 'Armory Tier', 'Rarities', 'Identified By',
                                                 'Recording', 'Seconds'], src_rows,
                   note='Every weapon in each weapon talent\'s Source panel, in the order the panel lists them (grouped by '
                        'class). Every weapon listed comes in all three rarities.')

    # ---- Weapon Talent Pools: the other way round, one row per weapon
    pool = {}
    for t in sources:
        for w in t['weapons']:
            pool.setdefault(w['name'], []).append(t['talent'])
    pr = []
    std = [a for a in arm_w if a['exotic'] != 'yes' and a['exotic'] is not True]
    cls_order = ['Assault Rifle', 'SMG', 'LMG', 'Marksman Rifle', 'Rifle', 'Shotgun']
    for a in sorted(std, key=lambda a: (cls_order.index(a['class']), a['armoryTier'], a['name'])):
        pr.append([a['name'], a['class'], a['armoryTier'], a['damageType'], None] + sorted(pool.get(a['name'], [])))
    for e in exotics:
        pr.append([e['name'], e['class'], 'T1', e['damageType'], len(e['possibleTalents'])] + sorted(e['possibleTalents']))
    heads = ['Weapon', 'Class', 'Armory Tier', 'Damage Type', 'Talents In Pool'] + [f'Talent {k}' for k in range(1, 19)]
    ws, r0 = sheet(wb, 'Weapon Talent Pools', heads, pr,
                   note='Each weapon\'s talent pool: the weapon talents it can roll (2 per weapon, randomly assigned from the '
                        'pool, per the database). Standard weapons come from the Source panels; exotics from their cards\' '
                        '"Possible Talents" grid, and an exotic also has its own fixed talent.')
    for i in range(r0 + 1, r0 + 1 + len(std)):
        ws.cell(row=i, column=5, value=f"=COUNTIF('Weapon Talent Sources'!$B:$B,A{i})").font = BODY_FONT

    # ---- Gear Set Effects
    gse = []
    for r in rows_csv('gear-set-effects.csv'):
        gse.append([r['Name'], int(r['List Order']), r['Icon'], r['2-Piece'], r['3-Piece'], r['4-Piece'], r['Armory Tiers'],
                    r['Shares Pieces With'], None, r['Recording'], secs(r)])
    ws, r0 = sheet(wb, 'Gear Set Effects', ['Name', 'List Order', 'Icon', '2-Piece', '3-Piece', '4-Piece', 'Armory Tiers',
                                            'Shares Pieces With', 'Pieces That Can Roll It', 'Recording', 'Seconds'], gse,
                   widths={'2-Piece': 34, '3-Piece': 40, '4-Piece': 60})
    for i in range(r0 + 1, r0 + 1 + len(gse)):
        ws.cell(row=i, column=9, value=f"=COUNTIFS('Gear Set Sources'!$A:$A,A{i},'Gear Set Sources'!$G:$G,\"<>*second listing*\")").font = BODY_FONT

    arm_fam = {a['name']: a['gearSetFamily'] for a in arm_g}

    def gear_src_rows(data, extra=True):
        out = []
        for t in data:
            for p in t['pieces']:
                out.append([t['entry'], p['name'], p['slot'], p['armoryTier'], p['family'], ' / '.join(p['rarities']),
                            'gear card' if p.get('cardCheck') else ('position' if p['match'] == 'position' else 'icon'),
                            t['source']['recording'], f"{t['source']['seconds'][0]} to {t['source']['seconds'][1]} s"])
            for x in (t.get('extraListings') or []) if extra else []:
                out.append([t['entry'], x.get('card', '(not identified yet)'), x['slot'], x['armoryTier'],
                            arm_fam.get(x.get('piece'), ''), x['rarity'],
                            'gear card (a second listing in this block)' if x.get('card') else 'not identified yet',
                            t['source']['recording'], f"{t['source']['seconds'][0]} to {t['source']['seconds'][1]} s"])
        return out

    heads_g = ['Entry', 'Piece', 'Slot', 'Armory Tier', 'Set Family', 'Rarities', 'Identified By', 'Recording', 'Seconds']
    gs_rows = gear_src_rows(gs_src)
    ws, r0 = sheet(wb, 'Gear Set Sources', heads_g, gs_rows,
                   note='Every gear piece in each Gear Set Effect\'s Source panel. "Position": named by its place in the list '
                        '(High-End then Superior; T2, T1, T0; mask, body armor, knee pads, backpack, gloves, holster), which '
                        'matches the Armory. "Gear card": the card opened from that tile shows the same name.')
    ws.cell(row=r0, column=1).value = 'Gear Set'
    ws.cell(row=1, column=1).value = (ws.cell(row=1, column=1).value + ' Captain Cow and Fearless Warrior list two pieces twice '
                                      'in their Superior T0 block; those rows say "a second listing".')

    # ---- Body Armor / Backpack talents and their sources
    for title, csv_name, data, src_title, slot in (
            ('Body Armor Talents', 'body-armor-talents.csv', ba_src, 'Body Armor Sources', 'Body Armor'),
            ('Backpack Talents', 'backpack-talents.csv', bp_src, 'Backpack Sources', 'Backpack')):
        rows = [[r['Name'], int(r['List Order']), r['Talent'], r['Talent Cooldown'], r['Shown At'], None, r['Note'],
                 r['Recording'], secs(r)] for r in rows_csv(csv_name)]
        ws, r0 = sheet(wb, title, ['Name', 'List Order', 'Talent', 'Talent Cooldown', 'Shown At',
                                   f'{slot} Pieces That Can Roll It', 'Note', 'Recording', 'Seconds'], rows)
        for i in range(r0 + 1, r0 + 1 + len(rows)):
            ws.cell(row=i, column=6, value=f"=COUNTIF('{src_title}'!$A:$A,A{i})").font = BODY_FONT
        ws2, r02 = sheet(wb, src_title, heads_g, gear_src_rows(data),
                         note=f'Every {slot.lower()} piece in each {slot} Talent\'s Source panel (named by matching its picture '
                              'with the gear set lists). Every piece listed comes in High-End and Superior.')
        ws2.cell(row=r02, column=1).value = 'Talent'

    # ---- Gear Pools: one row per Armory piece
    tal_of = {}
    for data in (ba_src, bp_src):
        for t in data:
            for p in t['pieces']:
                tal_of.setdefault(p['name'], []).append(t['entry'])
    slots = ['Mask', 'Body Armor', 'Knee Pads', 'Backpack', 'Gloves', 'Holster']
    gp = []
    for a in sorted(arm_g, key=lambda a: (slots.index(a['slot']), a['armoryTier'], a['gearSetFamily'])):
        gp.append([a['name'], a['slot'], a['armoryTier'], a['gearSetFamily'], ', '.join(a['gearSets']), None]
                  + sorted(tal_of.get(a['name'], [])))
    heads = ['Piece', 'Slot', 'Armory Tier', 'Set Family', 'Gear Sets It Can Roll', 'Talents In Pool'] + \
            [f'Talent {k}' for k in range(1, 13)]
    ws, r0 = sheet(wb, 'Gear Pools', heads, gp,
                   note='Each Armory gear piece: the gear sets it can carry and, for body armor and backpacks, the talents it '
                        'can roll. Gloves, holsters, knee pads and masks have no talent ("Generally, gear does not have a '
                        'Talent, except for Backpacks and Body Armor").')
    for i, row in enumerate(gp, r0 + 1):
        if row[1] == 'Body Armor':
            f = f"=COUNTIF('Body Armor Sources'!$B:$B,A{i})"
        elif row[1] == 'Backpack':
            f = f"=COUNTIF('Backpack Sources'!$B:$B,A{i})"
        else:
            f = 'none'
        ws.cell(row=i, column=6, value=f).font = BODY_FONT

    # ---- Bonus attributes, Armory
    ba = [[r['Name'], int(r['List Order']), r['Icon Group'], r['Description'], r['Current (Cory, at recording)'], r['Note'],
           r['Recording'], secs(r)] for r in rows_csv('bonus-attributes.csv')]
    sheet(wb, 'Bonus Attributes', ['Name', 'List Order', 'Icon Group', 'Description', "Current (Cory's character)", 'Note',
                                   'Recording', 'Seconds'], ba)
    aw = [[r['Name'], r['Class'], r['Exotic'], r['Armory Tier'], r['Season'], r['Damage Type'], r['Recording'], secs(r)]
          for r in rows_csv('armory-weapons.csv')]
    sheet(wb, 'Armory Weapons', ['Name', 'Class', 'Exotic', 'Armory Tier', 'Season', 'Damage Type', 'Recording', 'Seconds'], aw)
    ag = [[r['Name'], r['Slot'], r['Armory Tier'], r['Season'], r['Gear Sets'], r['Set Family'], r['Gear Sets Note'],
           r['Recording'], secs(r)] for r in rows_csv('armory-gear.csv')]
    sheet(wb, 'Armory Gear', ['Name', 'Slot', 'Armory Tier', 'Season', 'Gear Sets', 'Set Family', 'Gear Sets Note',
                              'Recording', 'Seconds'], ag)

    # ---- What the database cards don't show: creator readings and the recalibration fix
    cr = [[r['Topic'], r['Item'], r['Reading'], r['Creator'], r['Video Date'], r['Timestamp'], r['URL'], r['Note']]
          for r in rows_csv('creator-readings.csv')]
    sheet(wb, 'Creator Readings', ['Topic', 'Item', 'Reading', 'Creator', 'Video Date', 'Timestamp', 'URL', 'Note'], cr,
          widths={'URL': 44})
    rf = []
    for fx in load('recalibration-fixes.json')['fixes']:
        seen = {s['weapon']: s for s in fx.get('seenInGame', [])}
        for w in fx['weapons']:
            s = seen.get(w['name'], {})
            rf.append([w['name'], fx['patch'], fx['date'], ', '.join(s.get('outsidePool', [])), s.get('where', ''),
                       s.get('note', '')])
    sheet(wb, 'Recalibration Fixes', ['Weapon', 'Fixed In', 'Date', 'Seen Outside Its Pool', 'Where', 'Note'], rf,
          widths={'Where': 50})

    # ---- Read Me
    r = readme
    r.column_dimensions['A'].width = 28
    r.column_dimensions['B'].width = 12
    r.column_dimensions['C'].width = 100
    r['A1'] = 'The Division Resurgence: Build-Making Database'
    r['A1'].font = TITLE
    r['A2'] = ('Transcribed from screen recordings of the game (Season 2 Phase 2) on October 28 and 29, 2026, by '
               'Resurgence Builds. Values are exactly what the in-game cards show, including the game\'s own wording.')
    r['A2'].font = BODY_FONT
    r['A4'], r['B4'], r['C4'] = 'Tab', 'Rows', 'What it is'
    for c in ('A4', 'B4', 'C4'):
        r[c].font, r[c].fill = HEAD_FONT, HEAD_FILL
    tabs = [
        ('OS Protocols', 'All 66 OS Protocols (Engineering, Firepower and Toughness tabs) at level 1 with 0/5 talent upgrades.'),
        ('Exotic Weapons', 'The 6 exotics at Level 1, T1, no mods, with the 18 talents each can roll ("Possible Talents").'),
        ('Skill Mod Combo Sets', 'All 90 sets, by specialization tab. The 2-piece bonus is a range; the 3-piece talent is shown at the lowest tier.'),
        ('Weapon Talents', 'All 73 weapon talents (67 that roll on standard weapons, 6 exotic talents) at rank I, with how many weapons can roll each.'),
        ('Weapon Talent Sources', 'Every weapon each talent can roll on, from the talents\' Source panels.'),
        ('Weapon Talent Pools', 'The same data by weapon: each weapon\'s 18 possible talents.'),
        ('Gear Set Effects', 'All 16 gear sets with their 2, 3 and 4-piece bonuses.'),
        ('Gear Set Sources', 'Every gear piece each set can roll on.'),
        ('Body Armor Talents', 'All 25 body armor talents at rank I.'),
        ('Body Armor Sources', 'Every body armor piece each body armor talent can roll on.'),
        ('Backpack Talents', 'All 25 backpack talents at rank I.'),
        ('Backpack Sources', 'Every backpack each backpack talent can roll on.'),
        ('Gear Pools', 'The same data by gear piece: the sets each piece carries and the 12 talents each body armor and backpack can roll.'),
        ('Bonus Attributes', 'The 24 Weapon/Gear Bonus Attributes with the card\'s description.'),
        ('Armory Weapons', 'Every weapon in the Armory (86 standard, 6 exotic) with its tier page and damage type.'),
        ('Armory Gear', 'Every gear piece in the Armory (72) with its tier page and the sets it carries.'),
        ('Creator Readings', 'On-screen values from creator videos and the owner\'s screenshots that the database cards don\'t '
                             'show: rank II values, upgraded and level 50 OS Protocols, inventory counters and more, with the '
                             'video and timestamp.'),
        ('Recalibration Fixes', 'Weapons whose recalibration pulled talents from the wrong pool until a patch fixed it, so older '
                                'copies can carry a talent outside their pool.'),
    ]
    for i, (tab, what) in enumerate(tabs, 5):
        r.cell(row=i, column=1, value=tab).font = BOLD
        hdr = 2 if tab in ('Weapon Talent Sources', 'Weapon Talent Pools', 'Gear Set Sources', 'Body Armor Sources',
                           'Backpack Sources', 'Gear Pools') else 1
        r.cell(row=i, column=2, value=f"=COUNTA('{tab}'!$A:$A)-{hdr}").font = BODY_FONT
        c = r.cell(row=i, column=3, value=what)
        c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical='top')
    n = 5 + len(tabs) + 1
    notes = [
        ('How to read it', None),
        ('Rarities', 'High-End (gold), Superior (purple), Specialized (blue). Every weapon in a talent\'s list comes in all three; '
                     'every gear piece in a set or talent list comes in High-End and Superior.'),
        ('Armory Tier', 'The Armory page an item sits on (Season 1, Tier 0, 1 or 2).'),
        ('Talent pools', 'Each standard weapon can roll 18 of the 67 weapon talents (each exotic also 18, plus its own); each body armor '
                         'and each backpack can roll 12 of its 25 talents. A weapon gets 2 talents from its pool; body armor and '
                         'backpacks get one.'),
        ('Gear set families', 'The 72 pieces form four families (A to D), one piece per slot and tier. A: Fury Strike, Quick Draw, One Shot '
                              'One Kill (T1 and up), Gunny Johnny (T2). B: Jackpot, Long-term Effect, Dr. Medic (T1+), Mechanical Enemy '
                              '(T2). C: Phalanx Attack, Boom-Shakalaka, Mechanical Expert (T1+), Fire Cycle (T2). D: Captain Cow, '
                              'Fearless Warrior, Self-Propelled Shield (T1+), Healing Elites (T2).'),
        ('Names', 'As the game writes them. The site calls one talent "Precise Support" (item tooltips); the database card says '
                  '"Precision Support".'),
        ('Listed twice', 'Captain Cow and Fearless Warrior show eight Superior T0 tiles, not six. The extra two open "Tay" Protective '
                         'Pads v2 (the same card as the first) and "Fearless Hunter\'s Code Gloves Level 1", a finished Level 1 item '
                         '(Health 722, Armor 60, +8 to each core attribute). They are second listings of pieces already in the list.'),
        ('Gear cards', 'Tapping a tile opens the gear piece\'s card with its name. 37 of the 72 pieces were read this way and every '
                       'one matched the name its position in the list gives it.'),
        ('Not in the database', 'Ammo Refill\'s card still shows its pre-Season 2 value (29% chance of 1 ammo); Ubisoft\'s Season 2 '
                                'Phase 1 notes changed it. The database can lag behind a patch.'),
    ]
    for i, (a, b) in enumerate(notes, n):
        r.cell(row=i, column=1, value=a).font = BOLD if b else TITLE
        if b:
            c = r.cell(row=i, column=3, value=b)
            c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical='top')
    wb.save(out)
    print(out)


if __name__ == '__main__':
    main(sys.argv[1])
